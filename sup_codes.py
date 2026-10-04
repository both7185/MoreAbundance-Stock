"""
Sup no. codes ("Sup no.xlsx"): what each part of a code means + the list of shops and their codes.

    M - 2 11
    |   | `- positions 3-4: running number (assigned automatically for a new shop)
    |   `--- position 2 : group   (0 MK,BTG / 1 Company / 2 ของสด / 3 ผัก / ... / 9 อื่นๆ)
    `------- position 1 : letter  (A MK,BTG / C Company / M Market / S Stock / X Other)

* Stored in sup_codes.json next to the app.
* The order of `letters` is the sort order of the daily Excel (A -> C -> M -> S -> X by default).
* Running numbers are shared by all letters with the same group digit
  (group 1: C-101, C-102, C-103, S-104 ... S-107), so the next number looks at every code in the group.
* A product's Sup no. is its ร้านหลัก's code (like the XLOOKUP in the original Data sheet).
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from products import clean, format_sup, parse_sup

SUP_PATH = Path(__file__).with_name("sup_codes.json")

DEFAULT = {
    "letters": [
        {"code": "A", "meaning": "MK,BTG"}, {"code": "C", "meaning": "Company"},
        {"code": "M", "meaning": "Market"}, {"code": "S", "meaning": "Stock"}, {"code": "X", "meaning": "Other"},
    ],
    "groups": [
        {"code": "0", "meaning": "MK,BTG"}, {"code": "1", "meaning": "Company"}, {"code": "2", "meaning": "ของสด"},
        {"code": "3", "meaning": "ผัก"}, {"code": "4", "meaning": "ผลไม้"}, {"code": "5", "meaning": "ของหวาน"},
        {"code": "6", "meaning": "ของแห้ง"}, {"code": "9", "meaning": "อื่นๆ"},
    ],
    "shops": [],
}


# --------------------------------------------------------------------------- #
# Load / save
# --------------------------------------------------------------------------- #
def normalize_sup(data: dict) -> dict:
    """Trim text, drop empty rows, upper-case letters, format shop codes as A-001."""
    def defs(rows, kind):
        out, seen = [], set()
        for r in rows or []:
            code = clean(r.get("code")).upper()
            if kind == "group":
                code = code[:1]
            if not code or code in seen:
                continue
            seen.add(code)
            out.append({"code": code, "meaning": clean(r.get("meaning"))})
        return out

    letters = defs(data.get("letters"), "letter")
    shops = []
    for r in data.get("shops") or []:
        name = clean(r.get("name"))
        if name:
            shops.append({"name": name, "code": format_sup(r.get("code"))})
    # keep the shop list ordered by code: position-1 order, then number (stable for equal codes)
    order = {d["code"]: i for i, d in enumerate(letters)}

    def key(shop):
        p = parse_sup(shop["code"])
        return (0, order.get(p[0], len(order)), p[0], p[1]) if p else (1, 0, "", 0)

    shops.sort(key=key)
    return {"letters": letters, "groups": defs(data.get("groups"), "group"), "shops": shops}


def load_sup(path: Path = SUP_PATH) -> dict:
    if not path.exists():
        return normalize_sup(DEFAULT)
    with open(path, encoding="utf-8") as f:
        return normalize_sup(json.load(f))


def to_json(data: dict) -> str:
    return json.dumps(normalize_sup(data), ensure_ascii=False, indent=2) + "\n"


def save_sup_local(data: dict, path: Path = SUP_PATH) -> None:
    path.write_text(to_json(data), encoding="utf-8")


# --------------------------------------------------------------------------- #
# Labels for dropdowns
# --------------------------------------------------------------------------- #
def meanings(data: dict, kind: str) -> dict[str, str]:
    return {d["code"]: d["meaning"] for d in data[kind]}


def letter_label(data: dict, letter: str) -> str:
    m = meanings(data, "letters").get(letter, "")
    return f"{letter} ({m})" if m else letter


def group_label(data: dict, group: str) -> str:
    m = meanings(data, "groups").get(group, "")
    return f"{group} ({m})" if m else group


def code_meaning(data: dict, code: str) -> str:
    """'M-211' -> 'Market · ของสด'"""
    p = parse_sup(code)
    if not p:
        return ""
    letter, num = p
    parts = [meanings(data, "letters").get(letter, ""), meanings(data, "groups").get(str(num // 100), "")]
    return " · ".join(x for x in parts if x)


def shop_map(data: dict) -> dict[str, str]:
    return {s["name"]: s["code"] for s in data["shops"]}


def shop_label(data: dict, name: str) -> str:
    """'ร้านลูกชิ้น' -> 'ร้านลูกชิ้น — M-211 (Market · ของสด)'"""
    code = shop_map(data).get(name)
    if not code:
        return name
    m = code_meaning(data, code)
    return f"{name} — {code}" + (f" ({m})" if m else "")


def letter_order(data: dict) -> list[str]:
    return [d["code"] for d in data["letters"]]


# --------------------------------------------------------------------------- #
# Next running number
# --------------------------------------------------------------------------- #
def next_code(data: dict, letter: str, group: str, also_used: list[str] = ()) -> str:
    """
    Next free code for letter + group digit, e.g. ('M', '2') -> 'M-212'.
    Positions 3-4 = highest number already used in that group digit (any letter) + 1;
    if 99 is taken, the smallest unused number instead.
    """
    g = int(group)
    used = set()
    for code in [s["code"] for s in data["shops"]] + list(also_used):
        p = parse_sup(code)
        if p and p[1] // 100 == g:
            used.add(p[1] % 100)
    n = (max(used) + 1) if used else 1
    if n > 99:
        free = [i for i in range(1, 100) if i not in used]
        if not free:
            raise ValueError(f"กลุ่ม {group} ใช้เลขครบ 99 แล้ว")
        n = free[0]
    return f"{letter.upper()}-{g}{n:02d}"


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def validate_sup(data: dict, products: list[dict] | None = None) -> tuple[list[str], list[str]]:
    errors, warnings = [], []
    letters, groups = set(meanings(data, "letters")), set(meanings(data, "groups"))
    for d in data["letters"]:
        if not re.fullmatch(r"[A-Z]", d["code"]):
            errors.append(f"ตำแหน่งที่ 1 '{d['code']}' ต้องเป็นตัวอักษร A-Z 1 ตัว")
    for d in data["groups"]:
        if not re.fullmatch(r"\d", d["code"]):
            errors.append(f"ตำแหน่งที่ 2 '{d['code']}' ต้องเป็นตัวเลข 0-9 1 ตัว")
    seen = set()
    for s in data["shops"]:
        if s["name"] in seen:
            errors.append(f"ชื่อร้าน '{s['name']}' ซ้ำ")
        seen.add(s["name"])
        p = parse_sup(s["code"])
        if not p:
            errors.append(f"'{s['name']}': Sup no. '{s['code']}' ต้องเป็นรูปแบบ A-001")
            continue
        if p[0] not in letters:
            warnings.append(f"'{s['name']}' ({s['code']}): ตำแหน่งที่ 1 '{p[0]}' ยังไม่มีในตารางความหมาย")
        if str(p[1] // 100) not in groups:
            warnings.append(f"'{s['name']}' ({s['code']}): ตำแหน่งที่ 2 '{p[1] // 100}' ยังไม่มีในตารางความหมาย")
    if products is not None:
        used = {p.get("ร้านหลัก") for p in products if p.get("ร้านหลัก")}
        missing = sorted(used - seen)
        if missing:
            errors.append("ลบร้านที่ยังมีสินค้าใช้อยู่ไม่ได้: " + ", ".join(missing))
    return errors, warnings


# --------------------------------------------------------------------------- #
# Excel import / export (layout of "Sup no.xlsx")
# --------------------------------------------------------------------------- #
def import_sup_excel(data: bytes) -> tuple[dict, list[str]]:
    """
    Reads:
      * code meanings: pairs like  A | A-MK,BTG   and   0 | 0-MK,BTG   (any sheet)
      * shops: pairs like  ร้านลูกชิ้น | M-211     (any sheet)
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    letters, groups, shops = [], [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [clean(c) for c in row]
            for a, b in zip(cells, cells[1:]):
                if not a or not b:
                    continue
                if re.fullmatch(r"[A-Za-z]", a) and b.upper().startswith(a.upper() + "-") and not parse_sup(b):
                    letters.append({"code": a.upper(), "meaning": b.split("-", 1)[1].strip()})
                elif re.fullmatch(r"\d", a) and b.startswith(a + "-"):
                    groups.append({"code": a, "meaning": b.split("-", 1)[1].strip()})
                elif parse_sup(b) and not parse_sup(a) and not re.fullmatch(r"[A-Za-z]|\d", a):
                    shops.append({"name": a, "code": b})
    notes = []
    if not shops:
        raise ValueError("ไม่พบรายชื่อร้านคู่กับรหัส (เช่น 'ร้านลูกชิ้น | M-211')")
    if not letters or not groups:
        notes.append("ไม่พบตารางความหมายรหัส — ใช้ความหมายเดิม")
    current = load_sup()
    return normalize_sup({
        "letters": letters or current["letters"],
        "groups": groups or current["groups"],
        "shops": shops,
    }), notes


def export_sup_excel(data: dict) -> bytes:
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    body = Font(name="Calibri", size=11)

    def header(ws, row, col, text):
        c = ws.cell(row, col, text)
        c.font, c.fill, c.border, c.alignment = head_font, head_fill, border, Alignment(horizontal="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "การตั้งชื่อ"
    header(ws, 1, 1, "ตำแหน่ง 1"); header(ws, 1, 2, "ความหมาย")
    header(ws, 1, 4, "ตำแหน่ง 2"); header(ws, 1, 5, "ความหมาย")
    for i, d in enumerate(data["letters"], start=2):
        for c, v in ((1, d["code"]), (2, f"{d['code']}-{d['meaning']}")):
            cell = ws.cell(i, c, v); cell.font, cell.border = body, border
    for i, d in enumerate(data["groups"], start=2):
        for c, v in ((4, int(d["code"])), (5, f"{d['code']}-{d['meaning']}")):
            cell = ws.cell(i, c, v); cell.font, cell.border = body, border
    for col, w in zip("ABCDE", [10, 18, 3, 10, 18]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("ชื่อ Sup")
    for c, h in enumerate(["ร้าน", "Sup no.", "ความหมาย"], start=1):
        header(ws2, 1, c, h)
    for i, s in enumerate(data["shops"], start=2):
        for c, v in enumerate([s["name"], s["code"], code_meaning(data, s["code"])], start=1):
            cell = ws2.cell(i, c, v); cell.font, cell.border = body, border
    for col, w in zip("ABC", [26, 10, 22]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
