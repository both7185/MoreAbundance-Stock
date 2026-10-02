"""
Monthly summary: one row per PO file -> (ชื่อ PO, ยอดรวม), grand total at the bottom.

Ordering rules (from the file name):
  1. "(อัพเดทใหม่)" / "(อัปเดตใหม่)" / "- อัพเดทใหม่" is ignored (and dropped from the shown name)
  2. sort by the date in the file name
  3. same date: no bracket -> (มื้อเย็น) -> (กล่องโฟม) -> anything else (alphabetical)
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

from po_extractor import THAI_MONTHS_ABBR, THAI_MONTHS_FULL, PurchaseOrder, parse_thai_date
from report_builder import (
    BORDER_ALL, CENTER, F_BODY, F_BOLD, F_HEAD, F_SUBTITLE, F_TITLE,
    FILL_GROUP, FILL_HEADER, NF_2DP, THIN,
)

# "อัพเดทใหม่", "อัพเดตใหม่", "อัปเดตใหม่", "อัปเดทใหม่", "update" ...
_UPDATE_RE = re.compile(r"\s*-?\s*(อั[พป]เด[ทต](ใหม่)?|update[d]?)\s*-?\s*", re.IGNORECASE)

# order for the bracket note when several POs share a date
SUFFIX_ORDER = ["", "มื้อเย็น", "กล่องโฟม"]


def clean_po_name(filename: str) -> str:
    """'PO 14 ก.ย. 69 (มื้อเย็น - อัพเดทใหม่).pdf' -> 'PO 14 ก.ย. 69 (มื้อเย็น)'"""
    stem = re.sub(r"\.pdf$", "", filename.strip(), flags=re.IGNORECASE)

    def fix_bracket(m: re.Match) -> str:
        inner = _UPDATE_RE.sub(" ", m.group(1)).strip(" -")
        return f" ({inner})" if inner else ""

    name = re.sub(r"\s*\(([^)]*)\)", fix_bracket, stem)
    return re.sub(r"\s+", " ", name).strip()


def bracket_note(clean_name: str) -> str:
    m = re.search(r"\(([^)]*)\)", clean_name)
    return m.group(1).strip() if m else ""


def date_from_name(name: str) -> date | None:
    """'PO 11 ก.ย. 69 (มื้อเย็น)' -> date(2026, 9, 11)"""
    no_brackets = re.sub(r"\([^)]*\)", " ", name)
    return parse_thai_date(no_brackets)


def _suffix_rank(note: str) -> tuple[int, str]:
    return (SUFFIX_ORDER.index(note), "") if note in SUFFIX_ORDER else (len(SUFFIX_ORDER), note)


@dataclass
class MonthlyRow:
    name: str            # shown in the Excel ('PO 11 ก.ย. 69 (มื้อเย็น)')
    filename: str
    day: date | None     # date used for sorting
    amount: float        # ยอดรวมของ PO
    printed_total: float | None
    computed_total: float
    date_source: str     # 'ชื่อไฟล์' | 'ใช้ในวันที่ (PDF)' | '-'


def status_text(r: "MonthlyRow") -> str:
    """Same text in the web table and the Excel file."""
    if r.printed_total is None:
        return "⚠️ ไม่พบยอด 'เป็นเงิน' ในใบสั่ง"
    diff = Decimal(str(r.computed_total)) - Decimal(str(r.printed_total))
    if abs(diff) < Decimal("0.005"):
        return "✅ ตรงกัน"
    return f"⚠️ ต่าง {float(diff):+,.2f}"


def build_rows(pos: list[PurchaseOrder], use_printed_total: bool = True) -> list[MonthlyRow]:
    rows = []
    for po in pos:
        name = clean_po_name(po.filename)
        d, src = date_from_name(name), "ชื่อไฟล์"
        if d is None:
            d, src = po.delivery_date, "ใช้ในวันที่ (PDF)"
        if d is None:
            src = "-"
        amount = po.stated_total if (use_printed_total and po.stated_total is not None) else po.computed_total
        rows.append(MonthlyRow(name, po.filename, d, amount, po.stated_total, po.computed_total, src))

    rows.sort(key=lambda r: (r.day is None, r.day or date.max, _suffix_rank(bracket_note(r.name)), r.name))
    return rows


def month_label(rows: list[MonthlyRow]) -> str:
    """'เดือนกันยายน 2569' (or a range if files span months)."""
    months = sorted({(r.day.year, r.day.month) for r in rows if r.day})
    if not months:
        return ""
    fmt = lambda ym: f"{THAI_MONTHS_FULL[ym[1] - 1]} {ym[0] + 543}"
    if len(months) == 1:
        return f"เดือน{fmt(months[0])}"
    return f"{fmt(months[0])} - {fmt(months[-1])}"


def month_short(rows: list[MonthlyRow]) -> str:
    """'ก.ย. 69' for file names."""
    months = sorted({(r.day.year, r.day.month) for r in rows if r.day})
    if not months:
        return ""
    fmt = lambda ym: f"{THAI_MONTHS_ABBR[ym[1] - 1]} {(ym[0] + 543) % 100:02d}"
    return fmt(months[0]) if len(months) == 1 else f"{fmt(months[0])}-{fmt(months[-1])}"


# --------------------------------------------------------------------------- #
# Excel (same look as the daily summary)
# --------------------------------------------------------------------------- #
def build_monthly_workbook(rows: list[MonthlyRow]) -> bytes:
    """
    A ใบสั่งซื้อ | B ยอดรวม | C ผลรวมรายการ | D สถานะ
    last row: รวมทั้งเดือน (=SUM of B and C)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "รวมยอดทั้งเดือน"
    headers = ["ใบสั่งซื้อ", "ยอดรวม", "ผลรวมรายการ", "สถานะ"]
    widths = [38, 18, 18, 22]
    n = len(headers)
    last_col = get_column_letter(n)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = "ตารางสรุปยอดรวมใบสั่งซื้อวัตถุดิบประกอบอาหาร"
    ws["A1"].font = F_TITLE
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"สรุปยอดรวม{month_label(rows)} ({len(rows)} ใบสั่งซื้อ)"
    ws["A2"].font = F_SUBTITLE

    # row 4: group header (dark), row 5: column headers
    ws.merge_cells(f"A4:{last_col}4")
    ws["A4"] = month_label(rows) or "รวมยอดทั้งเดือน"
    ws["A4"].font, ws["A4"].fill, ws["A4"].alignment, ws["A4"].border = F_HEAD, FILL_GROUP, CENTER, BORDER_ALL
    for c in range(2, n + 1):
        ws.cell(4, c).border = Border(right=THIN if c == n else None, top=THIN, bottom=THIN)
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(5, c, text)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEADER, CENTER, BORDER_ALL

    first = 6
    for i, r in enumerate(rows):
        row = first + i
        values = [(r.name, None), (r.amount, NF_2DP), (r.computed_total, NF_2DP), (status_text(r), None)]
        for c, (v, nf) in enumerate(values, start=1):
            cell = ws.cell(row, c, v)
            cell.font, cell.border = F_BODY, BORDER_ALL
            if nf:
                cell.number_format = nf
    last = first + len(rows) - 1

    total_row = last + 1
    t = ws.cell(total_row, 1, "รวมทั้งเดือน")
    t.font, t.border, t.alignment = F_BOLD, BORDER_ALL, Alignment(horizontal="right")
    for c in (2, 3):
        col = get_column_letter(c)
        s = ws.cell(total_row, c, f"=SUM({col}{first}:{col}{last})")
        s.font, s.border, s.number_format = F_BOLD, BORDER_ALL, NF_2DP
    ws.cell(total_row, 4).border = BORDER_ALL

    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A6"
    # print: all 4 columns on one page width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def grand_total(rows: list[MonthlyRow]) -> float:
    return float(sum(Decimal(str(r.amount)) for r in rows))


# --------------------------------------------------------------------------- #
# Zip support (upload a whole month's folder as one .zip)
# --------------------------------------------------------------------------- #
def pdfs_from_zip(data: bytes) -> list[tuple[str, bytes]]:
    """Return [(file name, bytes)] for every PDF inside, with Thai names decoded."""
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            name = info.filename
            if not info.flag_bits & 0x800:  # name not flagged as UTF-8 -> try UTF-8 then Thai (cp874)
                raw = name.encode("cp437")
                for enc in ("utf-8", "cp874"):
                    try:
                        name = raw.decode(enc)
                        break
                    except UnicodeDecodeError:
                        continue
            base = name.replace("\\", "/").split("/")[-1]
            if base.lower().endswith(".pdf") and not base.startswith(("._", "~$")):
                out.append((base, z.read(info)))
    return out
