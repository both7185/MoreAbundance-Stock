"""
Aggregation + Excel generation for
'ตารางสรุปและรวมรายการใบสั่งซื้อวัตถุดิบประกอบอาหาร'
(one จำนวน column per PO date and tag, rows sorted by Sup no. from products.csv).

A PO whose file name has a bracket tag gets its own column right after that date's normal column:
  'PO 8 ต.ค. 69.pdf'                       -> 8 ต.ค. 69
  'PO 8 ต.ค. 69 (กล่องโฟม).pdf'             -> 8 ต.ค. 69 (กล่องโฟม)
  'PO 8 ต.ค. 69 (มื้อเย็น - อัพเดทใหม่).pdf' -> 8 ต.ค. 69 (มื้อเย็น)
"อัพเดทใหม่" on its own is not a tag, so 'PO 14 ก.ย. 69 (อัพเดทใหม่).pdf' stays in the normal column.
"""
from __future__ import annotations

import io
import json
import re
from decimal import Decimal
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from po_extractor import THAI_MONTHS_ABBR, PurchaseOrder, thai_short_date

CATALOG_PATH = Path(__file__).with_name("catalog.json")

# "อัพเดทใหม่", "อัพเดตใหม่", "อัปเดตใหม่", "อัปเดทใหม่", "update" ...
UPDATE_RE = re.compile(r"\s*-?\s*(อั[พป]เด[ทต](ใหม่)?|update[d]?)\s*-?\s*", re.IGNORECASE)

# column order within one date: normal, then a tag containing "โฟม", then "เย็น"; other tags follow, alphabetical
TAG_ORDER = ["โฟม", "เย็น"]


def po_tag(filename: str) -> str:
    """Column tag from the file name's brackets ('' = normal column). '(2)' copy numbers are ignored."""
    stem = re.sub(r"\.pdf$", "", filename.strip(), flags=re.IGNORECASE)
    parts = []
    for inner in re.findall(r"\(([^)]*)\)", stem):
        t = re.sub(r"\s+", " ", UPDATE_RE.sub(" ", inner)).strip(" -")
        if t and not t.isdigit():
            parts.append(t)
    return " ".join(parts)


def is_update(filename: str) -> bool:
    return bool(UPDATE_RE.search(filename))


Col = tuple  # (date, tag): one จำนวน column in the Excel


def col_sort_key(col: Col):
    d, tag = col
    if not tag:
        return (d, -1, "")
    rank = next((i for i, word in enumerate(TAG_ORDER) if word in tag), len(TAG_ORDER))
    return (d, rank, tag)


def col_label(col: Col, sep: str = " ") -> str:
    """(8 Oct, 'โฟม') -> '8 ต.ค. 69 (โฟม)'; sep="\\n" puts the tag on a second line (Excel header)"""
    d, tag = col
    return f"{thai_short_date(d)}{sep}({tag})" if tag else thai_short_date(d)


# --------------------------------------------------------------------------- #
# Catalog: only name / unit clean-up now (categories come from products.csv)
# --------------------------------------------------------------------------- #
def load_catalog(path: Path = CATALOG_PATH) -> dict:
    with open(path, encoding="utf-8") as f:
        catalog = json.load(f)
    for key in ("name_aliases", "unit_aliases"):
        m = catalog.get(key)
        catalog[key] = {k: v for k, v in (m or {}).items() if isinstance(k, str) and isinstance(v, str) and v.strip()}
    return catalog


# --------------------------------------------------------------------------- #
# Aggregation
# --------------------------------------------------------------------------- #
@dataclass
class SummaryRow:
    category: str           # หมวดหมู่ from products.csv ('' if not found)
    name: str
    unit: str
    unit_price: float
    sup_no: str = ""        # used for sorting only - not written to the Excel
    matched: bool = False   # found in products.csv
    product_name: str = ""  # name of the matched product in products.csv
    po_names: set = field(default_factory=set)   # names exactly as printed in the POs
    qty_by_col: dict[Col, float] = field(default_factory=dict)
    cost: float | None = None  # ทุน per unit from products.csv (None = unknown)


@dataclass
class Summary:
    columns: list[Col]                    # (date, tag), in Excel order
    rows: list[SummaryRow]
    po_total_by_col: dict[Col, float]     # "Original PO Total" per column
    po_numbers: list[str]
    files_by_col: dict[Col, list[str]] = field(default_factory=dict)

    @property
    def dates(self) -> list[date]:
        return sorted({d for d, _ in self.columns})

    @property
    def unmatched(self) -> list[SummaryRow]:
        return [r for r in self.rows if not r.matched]

    @property
    def missing_category(self) -> list[SummaryRow]:
        """In products.csv but without หมวดหมู่ (one row per product)."""
        seen, out = set(), []
        for r in self.rows:
            if r.matched and not r.category and r.product_name not in seen:
                seen.add(r.product_name)
                out.append(r)
        return out


def aggregate(
    pos: list[PurchaseOrder],
    date_for_po: dict[str, date],
    catalog: dict,
    product_lookup: dict | None = None,
    apply_aliases: bool = True,
) -> Summary:
    """
    Group by (item name, unit, unit price); split quantity by column (date + file-name tag).
    Order: Sup no. (A -> C -> M -> S -> X, then number), then the product's row in the product list;
    items not in the product list go last (by name).
    """
    from products import norm, sup_sort_key  # local import keeps this module usable on its own

    product_lookup = product_lookup or {}
    name_alias = catalog.get("name_aliases", {}) if apply_aliases else {}
    unit_alias = catalog.get("unit_aliases", {}) if apply_aliases else {}

    groups: "OrderedDict[tuple, SummaryRow]" = OrderedDict()
    po_totals: dict[Col, float] = {}
    files_by_col: dict[Col, list[str]] = {}
    po_numbers: list[str] = []
    order_of: dict[tuple, int] = {}

    for po in pos:
        d = (date_for_po[po.filename], po_tag(po.filename))
        files_by_col.setdefault(d, []).append(po.filename)
        if po.po_number and po.po_number not in po_numbers:
            po_numbers.append(po.po_number)
        total = po.stated_total if po.stated_total is not None else po.computed_total
        po_totals[d] = float(Decimal(str(po_totals.get(d, 0.0))) + Decimal(str(total)))

        for it in po.items:
            name = name_alias.get(it.name, it.name)
            unit = unit_alias.get(it.unit, it.unit)
            key = (name, unit, it.unit_price)
            row = groups.get(key)
            if row is None:
                info = product_lookup.get(norm(name)) or product_lookup.get(norm(it.name))
                row = groups[key] = SummaryRow(
                    category=info.category if info else "",
                    name=name, unit=unit, unit_price=it.unit_price,
                    sup_no=info.sup_no if info else "", matched=info is not None,
                    product_name=info.name if info else "",
                    cost=getattr(info, "cost", None) if info else None,
                )
                order_of[key] = info.order if info else 10**9
            row.po_names.add(it.name)
            row.qty_by_col[d] = float(Decimal(str(row.qty_by_col.get(d, 0.0))) + Decimal(str(it.qty)))

    rows = sorted(
        groups.items(),
        key=lambda kv: (
            not kv[1].matched,              # items not in the product list go last
            sup_sort_key(kv[1].sup_no),     # A-001, A-002, ... C-..., M-..., S-..., X-...
            order_of[kv[0]],                # same Sup no.: order of the product list
            kv[1].name, kv[1].unit, kv[1].unit_price,
        ),
    )
    columns = sorted(po_totals, key=col_sort_key)
    return Summary(columns, [r for _, r in rows], po_totals, po_numbers, files_by_col)


def date_range_label(dates: list[date]) -> str:
    """[6 Oct, 7 Oct] -> '6-7 ต.ค. 69'; mixed months -> '30 ก.ย. 69 - 2 ต.ค. 69'."""
    if not dates:
        return ""
    if len(dates) == 1:
        return thai_short_date(dates[0])
    same_month = all((d.year, d.month) == (dates[0].year, dates[0].month) for d in dates)
    contiguous = all((b - a).days == 1 for a, b in zip(dates, dates[1:]))
    if same_month:
        tail = f"{THAI_MONTHS_ABBR[dates[0].month - 1]} {(dates[0].year + 543) % 100:02d}"
        days = f"{dates[0].day}-{dates[-1].day}" if contiguous else ", ".join(str(d.day) for d in dates)
        return f"{days} {tail}"
    if contiguous:
        return f"{thai_short_date(dates[0])} - {thai_short_date(dates[-1])}"
    return ", ".join(thai_short_date(d) for d in dates)


# --------------------------------------------------------------------------- #
# Excel (styles copied from the template)
# --------------------------------------------------------------------------- #
FONT_NAME = "Calibri"
THIN = Side(style="thin")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_GROUP = PatternFill("solid", fgColor="1F497D")    # row 4
FILL_HEADER = PatternFill("solid", fgColor="2F5597")   # row 5
FILL_STRIPE = PatternFill("solid", fgColor="DDEBF7")   # every other body row (light blue / white)
F_TITLE = Font(name=FONT_NAME, size=14, bold=True)
F_SUBTITLE = Font(name=FONT_NAME, size=11)
F_HEAD = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT_NAME, size=11)
F_BOLD = Font(name=FONT_NAME, size=11, bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
NF_2DP = "#,##0.00"
NF_INT = "#,##0"

FIXED_COLS = [  # header, width
    ("ลำดับ", 6), ("หมวดหมู่", 12), ("รายการสินค้า", 25), ("หน่วย", 8),
    ("จำนวนรวม", 12), ("ราคา/หน่วย", 15),
]
DATE_COL_WIDTH = 10
PROFIT_COLS = [("ทุน", 11), ("ทุนรวม", 14), ("กำไร", 11), ("กำไรรวม", 14)]  # header, width
TOTAL_COL_WIDTH = 18


def _merged_header(ws, row, c1, c2, text, font, fill):
    """Header block (merged when it spans >1 column) with the template's border pattern."""
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    anchor = ws.cell(row, c1, text)
    anchor.font, anchor.fill, anchor.alignment, anchor.border = font, fill, CENTER, BORDER_ALL
    for c in range(c1 + 1, c2 + 1):
        ws.cell(row, c).border = Border(
            right=THIN if c == c2 else None, top=THIN, bottom=THIN)


def _merged_label(ws, row, c1, c2, text):
    """Footer label (merged, bold, right aligned)."""
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    a = ws.cell(row, c1, text)
    a.font, a.border = F_BOLD, BORDER_ALL
    a.alignment = Alignment(horizontal="right")
    for c in range(c1 + 1, c2 + 1):
        ws.cell(row, c).border = Border(right=THIN if c == c2 else None, top=THIN, bottom=THIN)


def _header_width(text: str) -> float:
    """Column width that fits a bold header; Thai vowel/tone marks sit above/below and take no width."""
    width = 0.0
    for ch in text:
        if ch == "ั" or "ิ" <= ch <= "ฺ" or "็" <= ch <= "๎":
            continue
        width += 1.4 if "฀" <= ch <= "๿" else 1.0   # Thai letters are wider than digits
    return max(DATE_COL_WIDTH, round(width + 2.5, 1))


def _nf(v: float) -> str:
    """#,##0.00 normally; extra decimals only if the value really has them (never hides digits)."""
    return NF_2DP if round(v, 2) == v else "#,##0.00####"


def _num(q: float):
    return int(q) if float(q).is_integer() else q


def build_workbook(summary: Summary, po_number_label: str | None = None) -> bytes:
    """
    Layout
      A ลำดับ | B หมวดหมู่ | C รายการสินค้า | D หน่วย | E จำนวนรวม | F ราคา/หน่วย
      G.. one 'จำนวน' column per date + tag  (e.g. '8 ต.ค. 69', '8 ต.ค. 69 (โฟม)')
      then ทุน | ทุนรวม (= ทุน × จำนวนรวม) | กำไร (= ราคา/หน่วย − ทุน) | กำไรรวม (= กำไร × จำนวนรวม)
      last 'จำนวนเงินรวม'               (= จำนวนรวม × ราคา/หน่วย)
    จำนวนรวม = SUM of the date columns, so editing any day's quantity updates everything.
    ทุน is blank when products.csv has no cost; the other three stay blank until a ทุน is typed.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "รวมใบสั่งซื้อ"
    L = get_column_letter

    n_fixed = len(FIXED_COLS)
    qty_col = 5                                   # E  จำนวนรวม
    price_col = 6                                 # F  ราคา/หน่วย
    first_date_col = n_fixed + 1                  # G
    last_date_col = first_date_col + len(summary.columns) - 1
    cost_col = last_date_col + 1                  # ทุน
    cost_total_col, profit_col, profit_total_col = cost_col + 1, cost_col + 2, cost_col + 3
    total_col = profit_total_col + 1
    last_col = total_col

    # ---- title rows -------------------------------------------------------
    po_label = po_number_label if po_number_label is not None else ", ".join(summary.po_numbers)
    subtitle = f"เปรียบเทียบและรวมยอดใบสั่งซื้อประจำวันที่ {date_range_label(summary.dates)}"
    if po_label:
        subtitle += f" (ใบสั่งที่ {po_label})"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, "ตารางสรุปและรวมรายการใบสั่งซื้อวัตถุดิบประกอบอาหาร").font = F_TITLE
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, subtitle).font = F_SUBTITLE

    # ---- header rows 4-5 --------------------------------------------------
    _merged_header(ws, 4, 1, n_fixed, "รายละเอียดสินค้า", F_HEAD, FILL_GROUP)
    for i, col in enumerate(summary.columns):
        _merged_header(ws, 4, first_date_col + i, first_date_col + i, col_label(col, "\n"), F_HEAD, FILL_GROUP)
        if col[1]:
            ws.cell(4, first_date_col + i).alignment = CENTER_WRAP
    _merged_header(ws, 4, cost_col, profit_total_col, "ทุน / กำไร", F_HEAD, FILL_GROUP)
    _merged_header(ws, 4, total_col, total_col, "รวมทั้งสิ้น", F_HEAD, FILL_GROUP)

    headers = ([h for h, _ in FIXED_COLS] + ["จำนวน"] * len(summary.columns)
               + [h for h, _ in PROFIT_COLS] + ["จำนวนเงินรวม"])
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(5, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEADER, CENTER, BORDER_ALL

    # ---- body -------------------------------------------------------------
    first_row = 6
    for idx, r in enumerate(summary.rows):
        row = first_row + idx
        for c, v in enumerate([idx + 1, r.category or None, r.name, r.unit], start=1):
            cell = ws.cell(row, c, v)
            cell.font, cell.border = F_BODY, BORDER_ALL

        total_qty = float(sum(Decimal(str(v)) for v in r.qty_by_col.values()))
        q = ws.cell(row, qty_col, f"=SUM({L(first_date_col)}{row}:{L(last_date_col)}{row})")
        q.font, q.border = F_BOLD, BORDER_ALL
        q.number_format = NF_INT if total_qty.is_integer() else _nf(total_qty)

        p = ws.cell(row, price_col, r.unit_price)
        p.font, p.border, p.number_format = F_BODY, BORDER_ALL, _nf(r.unit_price)

        for i, col in enumerate(summary.columns):
            cell = ws.cell(row, first_date_col + i)
            cell.font, cell.border = F_BODY, BORDER_ALL
            if col in r.qty_by_col:
                v = r.qty_by_col[col]
                cell.value = _num(v)
                cell.number_format = NF_INT if float(v).is_integer() else _nf(v)

        k = ws.cell(row, cost_col, r.cost)
        k.font, k.border = F_BODY, BORDER_ALL
        k.number_format = _nf(r.cost) if r.cost is not None else NF_2DP
        cost_ref, qty_ref = f"{L(cost_col)}{row}", f"{L(qty_col)}{row}"
        profit_cells = [
            (cost_total_col, f'=IF({cost_ref}="","",{cost_ref}*{qty_ref})'),
            (profit_col, f'=IF({cost_ref}="","",{L(price_col)}{row}-{cost_ref})'),
            (profit_total_col, f'=IF({cost_ref}="","",{L(profit_col)}{row}*{qty_ref})'),
        ]
        for c, formula in profit_cells:
            cell = ws.cell(row, c, formula)
            cell.font, cell.border, cell.number_format = F_BODY, BORDER_ALL, NF_2DP
        ws.cell(row, profit_total_col).font = F_BOLD

        t = ws.cell(row, total_col, f"={L(qty_col)}{row}*{L(price_col)}{row}")
        t.font, t.border, t.number_format = F_BOLD, BORDER_ALL, NF_2DP

    last_body = first_row + len(summary.rows) - 1
    for row in range(first_row + 1, last_body + 1, 2):   # 2nd, 4th, ... item rows
        for c in range(1, last_col + 1):
            ws.cell(row, c).fill = FILL_STRIPE
    r_total, r_orig, r_var = last_body + 1, last_body + 2, last_body + 3

    # ---- footer -----------------------------------------------------------
    _merged_label(ws, r_total, 1, qty_col - 1, "รวมทั้งสิ้น (Calculated Total)")
    _merged_label(ws, r_orig, 1, qty_col - 1, "ยอดรวมตามใบสั่งซื้อต้นฉบับ (Original PO Total)")
    _merged_label(ws, r_var, 1, qty_col - 1, "ผลต่าง / ตรวจสอบความถูกต้อง (Variance Check)")

    def footer(row, col, value, nf=None, center=False):
        cell = ws.cell(row, col, value)
        cell.font, cell.border = F_BOLD, BORDER_ALL
        if nf:
            cell.number_format = nf
        if center:
            cell.alignment = Alignment(horizontal="center")

    # quantity columns (จำนวนรวม + each date): sum only, no money comparison
    def qty_total_nf(values) -> str:
        t = float(sum(Decimal(str(v)) for v in values))
        return NF_INT if t.is_integer() else _nf(t)   # 2.5 kg must not show as 3

    col_values = {qty_col: [q for r in summary.rows for q in r.qty_by_col.values()]}
    for i, col in enumerate(summary.columns):
        col_values[first_date_col + i] = [r.qty_by_col[col] for r in summary.rows if col in r.qty_by_col]
    for c in [qty_col] + list(range(first_date_col, last_date_col + 1)):
        footer(r_total, c, f"=SUM({L(c)}{first_row}:{L(c)}{last_body})", qty_total_nf(col_values[c]))
        footer(r_orig, c, "-", center=True)
        footer(r_var, c, "-", center=True)
    for rr in (r_total, r_orig, r_var):
        for c in (price_col, cost_col, profit_col):
            footer(rr, c, None)
    for c in (cost_total_col, profit_total_col):   # SUM skips the blank ("") cells of items without ทุน
        footer(r_total, c, f"=SUM({L(c)}{first_row}:{L(c)}{last_body})", NF_2DP)
        footer(r_orig, c, "-", center=True)
        footer(r_var, c, "-", center=True)

    # money: grand total vs. sum of the printed PO totals
    orig_total = float(sum(Decimal(str(v)) for v in summary.po_total_by_col.values()))
    footer(r_total, total_col, f"=SUM({L(total_col)}{first_row}:{L(total_col)}{last_body})", NF_2DP)
    footer(r_orig, total_col, orig_total, NF_2DP)
    footer(r_var, total_col, f"={L(total_col)}{r_total}-{L(total_col)}{r_orig}", NF_2DP)

    # ---- layout -----------------------------------------------------------
    for c, (_, w) in enumerate(FIXED_COLS, start=1):
        ws.column_dimensions[L(c)].width = w
    for i, col in enumerate(summary.columns):
        ws.column_dimensions[L(first_date_col + i)].width = _header_width(f"({col[1]})") if col[1] else DATE_COL_WIDTH
    if any(tag for _, tag in summary.columns):
        ws.row_dimensions[4].height = 30   # date on top, (tag) below
    for c, (_, w) in enumerate(PROFIT_COLS, start=cost_col):
        ws.column_dimensions[L(c)].width = w
    ws.column_dimensions[L(total_col)].width = TOTAL_COL_WIDTH
    ws.freeze_panes = f"{L(first_date_col)}{first_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
