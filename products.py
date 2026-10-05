"""
Product master data ("Data P Both"): หมวดหมู่ + Sup no. for every product.

* Stored in products.csv (UTF-8) next to the app. One row per product.
* A PO item is matched to a product by name, ignoring spaces. Extra spellings used in POs can be
  listed in the 'ชื่อใน PO' column, separated by '|'.
* Sup no. comes from the product's ร้านหลัก (shop list in sup_codes.json), like the original XLOOKUP.
* The daily Excel is sorted by Sup no.: letter in LETTER_ORDER (A -> C -> M -> S -> X, set from
  sup_codes.json), then the number (A-001, A-002, ...), then the product's row order in this list.
* Saving on Streamlit Community Cloud: the disk there is temporary, so if GitHub secrets are set
  (see README) every save is also committed to the GitHub `data` branch (datastore.py).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PRODUCTS_PATH = Path(__file__).with_name("products.csv")

COL_NO = "ลำดับ"
COL_CAT = "หมวดหมู่"
COL_NAME = "รายการ"
COL_UNIT = "หน่วยนับ"
COL_SUPPLIER = "ร้านหลัก"
COL_SUP = "Sup no."
COL_ALIAS = "ชื่อใน PO"           # other spellings found in POs, separated by |
COLUMNS = [COL_NO, COL_CAT, COL_NAME, COL_UNIT, COL_SUPPLIER, COL_SUP, COL_ALIAS]

LETTER_ORDER = ["A", "C", "M", "S", "X"]
SUP_RE = re.compile(r"^\s*([A-Za-z]+)\s*-?\s*(\d+)\s*$")


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def norm(name) -> str:
    """Compare names ignoring spaces / case / Unicode form ('อกไก่ (ไม่หั่น)' == 'อกไก่(ไม่หั่น)')."""
    s = unicodedata.normalize("NFC", str(name or ""))
    s = s.replace("ํา", "ำ")  # ํา -> ำ
    return re.sub(r"\s+", "", s).casefold()


def clean(v) -> str:
    """Cell value -> trimmed text ('' for None / NaN)."""
    if v is None:
        return ""
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        if v.is_integer():
            v = int(v)
    return str(v).strip()


def split_aliases(text: str) -> list[str]:
    return [a.strip() for a in re.split(r"[|\n]", clean(text)) if a.strip()]


def parse_sup(code: str) -> tuple[str, int] | None:
    m = SUP_RE.match(clean(code))
    return (m.group(1).upper(), int(m.group(2))) if m else None


def format_sup(code: str) -> str:
    """'a-1' -> 'A-001'; anything unparseable is returned as typed."""
    p = parse_sup(code)
    return f"{p[0]}-{p[1]:03d}" if p else clean(code)


def set_letter_order(letters: list[str]) -> None:
    """Sort order of position 1 (from sup_codes.json)."""
    global LETTER_ORDER
    if letters:
        LETTER_ORDER = list(letters)


_LEADING_VOWELS = "\u0E40\u0E41\u0E42\u0E43\u0E44"  # เ แ โ ใ ไ


def thai_sort_key(text: str) -> str:
    """Thai dictionary order: 'เบทาโกร' sorts under บ, 'แมคโคร' under ม (not after ฮ)."""
    s = clean(text)
    s = re.sub(f"([{_LEADING_VOWELS}])([\u0E01-\u0E2E])", r"\2\1", s)
    return re.sub("[\u0E48-\u0E4C]", "", s)  # tone marks don't decide the order


def sup_sort_key(code: str) -> tuple:
    p = parse_sup(code)
    if not p:
        return (1, 0, "", 0)
    letter, num = p
    rank = LETTER_ORDER.index(letter) if letter in LETTER_ORDER else len(LETTER_ORDER)
    return (0, rank, letter, num)


# --------------------------------------------------------------------------- #
# Load / save
# --------------------------------------------------------------------------- #
def normalize_rows(rows: list[dict], shop_codes: dict[str, str] | None = None) -> list[dict]:
    """
    Trim text, drop empty rows, format Sup no., renumber ลำดับ.
    Sup no. = code of the ร้านหลัก in `shop_codes` (shop list); otherwise kept as typed,
    or filled from another product of the same shop.
    """
    out = []
    for r in rows:
        row = {c: clean(r.get(c)) for c in COLUMNS}
        if not row[COL_NAME]:
            continue
        row[COL_SUP] = format_sup(row[COL_SUP])
        row[COL_ALIAS] = " | ".join(split_aliases(row[COL_ALIAS]))
        out.append(row)
    if shop_codes:
        for row in out:
            if row[COL_SUPPLIER] in shop_codes:
                row[COL_SUP] = shop_codes[row[COL_SUPPLIER]]
    # same supplier -> same Sup no. (fallback for shops not in the shop list)
    by_supplier = {}
    for row in out:
        if row[COL_SUPPLIER] and parse_sup(row[COL_SUP]):
            by_supplier.setdefault(row[COL_SUPPLIER], row[COL_SUP])
    for row in out:
        if not row[COL_SUP] and row[COL_SUPPLIER] in by_supplier:
            row[COL_SUP] = by_supplier[row[COL_SUPPLIER]]
    for i, row in enumerate(out, start=1):
        row[COL_NO] = str(i)
    return out


def load_products(path: Path = PRODUCTS_PATH) -> list[dict]:
    if not path.exists():
        return []
    with open(path, encoding="utf-8-sig", newline="") as f:
        return normalize_rows(list(csv.DictReader(f)))


def to_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS, lineterminator="\n")
    w.writeheader()
    for r in rows:
        w.writerow({c: r.get(c, "") for c in COLUMNS})
    return buf.getvalue()


def save_products_local(rows: list[dict], path: Path = PRODUCTS_PATH) -> None:
    path.write_text(to_csv(rows), encoding="utf-8")


def validate(rows: list[dict]) -> tuple[list[str], list[str]]:
    """(errors that block saving, warnings)."""
    errors, warnings = [], []
    seen: dict[str, str] = {}
    for r in rows:
        for nm in [r[COL_NAME]] + split_aliases(r[COL_ALIAS]):
            k = norm(nm)
            if k in seen and seen[k] != r[COL_NAME]:
                errors.append(f"ชื่อ '{nm}' ซ้ำกับสินค้า '{seen[k]}'")
            elif k in seen and nm != r[COL_NAME]:
                pass
            elif k in seen:
                errors.append(f"สินค้า '{nm}' มีมากกว่า 1 แถว")
            seen.setdefault(k, r[COL_NAME])
        if r[COL_SUP] and not parse_sup(r[COL_SUP]):
            warnings.append(f"'{r[COL_NAME]}': Sup no. '{r[COL_SUP]}' ไม่อยู่ในรูปแบบ A-001 — จะถูกเรียงไว้ท้าย")
        if not r[COL_SUP]:
            warnings.append(f"'{r[COL_NAME]}': ยังไม่มี Sup no. — จะถูกเรียงไว้ท้าย")
    return errors, warnings


# --------------------------------------------------------------------------- #
# Lookup used by the report
# --------------------------------------------------------------------------- #
@dataclass
class ProductInfo:
    name: str
    category: str
    sup_no: str
    order: int  # row position in the list (tie-break inside the same Sup no.)


def build_lookup(rows: list[dict]) -> dict[str, ProductInfo]:
    lookup: dict[str, ProductInfo] = {}
    for i, r in enumerate(rows):
        info = ProductInfo(r[COL_NAME], r[COL_CAT], r[COL_SUP], i)
        for nm in [r[COL_NAME]] + split_aliases(r[COL_ALIAS]):
            lookup.setdefault(norm(nm), info)
    return lookup


def suggest(name: str, rows: list[dict], n: int = 1, cutoff: float = 0.0) -> list[tuple[str, float]]:
    """Closest product names for an unmatched PO item."""
    k = norm(name)
    scored = [(r[COL_NAME], SequenceMatcher(None, k, norm(r[COL_NAME])).ratio()) for r in rows]
    scored = [s for s in scored if s[1] >= cutoff]
    return sorted(scored, key=lambda s: -s[1])[:n]


def suggest_category(rows: list[dict], shop: str = "", sup_no: str = "") -> str:
    """Most common หมวดหมู่ among products of the same ร้านหลัก (else the same Sup no.)."""
    from collections import Counter
    for key, value in ((COL_SUPPLIER, shop), (COL_SUP, sup_no)):
        if value:
            c = Counter(r[COL_CAT] for r in rows if r.get(key) == value and r.get(COL_CAT))
            if c:
                return c.most_common(1)[0][0]
    return ""


def add_alias(rows: list[dict], product_name: str, alias: str) -> bool:
    for r in rows:
        if r[COL_NAME] == product_name:
            current = split_aliases(r[COL_ALIAS])
            if norm(alias) not in {norm(a) for a in current + [r[COL_NAME]]}:
                r[COL_ALIAS] = " | ".join(current + [alias])
            return True
    return False


# --------------------------------------------------------------------------- #
# Excel import / export
# --------------------------------------------------------------------------- #
_HEADER_KEYS = {
    COL_CAT: ["หมวดหมู่"], COL_NAME: ["รายการ", "รายการสินค้า", "ชื่อสินค้า"], COL_UNIT: ["หน่วยนับ", "หน่วย"],
    COL_SUPPLIER: ["ร้านหลัก", "ร้าน"], COL_SUP: ["sup no.", "sup no", "supno", "sup"], COL_ALIAS: ["ชื่อใน po"],
}


def import_excel(data: bytes) -> tuple[list[dict], list[str]]:
    """Read a 'Data P Both' style sheet. Uses the values Excel saved for formula cells."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    notes: list[str] = []
    for ws in wb.worksheets:
        values = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        for hi, header in enumerate(values[:20]):
            hdr = [clean(h).casefold() for h in header]
            colmap = {}
            for col, keys in _HEADER_KEYS.items():
                for i, h in enumerate(hdr):
                    if h in keys and col not in colmap:
                        colmap[col] = i
            if COL_NAME in colmap:
                rows, missing_code = [], 0
                for raw in values[hi + 1:]:
                    rec = {col: (raw[i] if i < len(raw) else None) for col, i in colmap.items()}
                    if not clean(rec.get(COL_NAME)):
                        continue
                    if COL_SUP in colmap and not clean(rec.get(COL_SUP)):
                        missing_code += 1
                    rows.append(rec)
                if missing_code:
                    notes.append(f"{missing_code} แถวไม่มีค่า Sup no. (ถ้าเป็นสูตร XLOOKUP ให้เปิดไฟล์ใน Excel แล้วกด Save ก่อน)")
                return normalize_rows(rows), notes
    raise ValueError("ไม่พบหัวตารางที่มีคอลัมน์ 'รายการ'")


def export_excel(rows: list[dict]) -> bytes:
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for c, h in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.border = head_font, head_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_i, r in enumerate(rows, start=2):
        for c, col in enumerate(COLUMNS, start=1):
            v = r.get(col, "")
            cell = ws.cell(r_i, c, int(v) if col == COL_NO and v.isdigit() else v)
            cell.font, cell.border = Font(name="Calibri", size=11), border
    for col, w in zip("ABCDEFG", [6, 14, 30, 9, 22, 9, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# GitHub settings for the permanent save on Streamlit Community Cloud (used by datastore.py)
# --------------------------------------------------------------------------- #
def github_config(secrets) -> dict | None:
    """Expects in .streamlit/secrets.toml (or the Cloud 'Secrets' box):
        [github]
        token  = "github_pat_..."     # fine-grained token, Contents: read & write on this repo
        repo   = "owner/repo"
        branch = "main"               # optional: code branch (the data branch starts from it)
        path   = "products.csv"       # optional
        data_branch = "data"          # optional: where the app saves products.csv / sup_codes.json
    """
    try:
        gh = dict(secrets["github"])
    except Exception:
        return None
    if not gh.get("token") or not gh.get("repo"):
        return None
    gh.setdefault("branch", "main")
    gh.setdefault("path", "products.csv")
    return gh

