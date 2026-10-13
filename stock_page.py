"""
หน้า 📦 สต็อกสินค้า (separate from the PO pages). Data logic lives in stock.py.

* Main table: click 📜 ดูประวัติ at the end of a row for that item's in/out history in a pop-up; click the
  วันหมดอายุ cell to edit its expiry lots (one lot per line). ✏️ แก้ไขตาราง switches the table to editing name / หมวด / หน่วย / ขั้นต่ำ (💾 บันทึก).
* 📜 ประวัติทั้งหมด: every item's movements in one pop-up.
* ➕ สินค้าเข้า: search an item, then เข้า / จำนวนในสต็อก / วันหมดอายุ / วันที่บันทึก → ยอดใหม่.
* 📤 เบิกสินค้าออก: pick the delivery date and place (one ใบเบิก per date + place, like a folder — open it
  again to edit), then
  add items to its list: item + expiry lot (the earliest expiry is picked first) + จำนวน. The stock is cut
  when the ใบเบิก is saved; editing it later changes the stock by the difference (stock.save_withdrawal).
  📥 Download Excel next to 💾 บันทึกใบเบิก gives the saved ใบเบิก of that day as a file (sheet_excel).
* 🆕 เพิ่มสินค้าใหม่: name, รหัส, หมวด only — quantities come from สินค้าเข้า.
* 🗑️ ลบสินค้า: search an item, tick to confirm, delete (its old movements stay in ประวัติทั้งหมด).
"""
from __future__ import annotations

import io
from datetime import date, timedelta

import pandas as pd
import streamlit as st

import stock
import storage
from stock import AVAILABLE, CATEGORIES, LOW, OUT, fmt_date, fmt_num

COLORS = {AVAILABLE: ("#1b7f3b", "#e3f4e8"), LOW: ("#c25e00", "#fdeedd"), OUT: ("#c62828", "#fde4e4")}
STATUS_TH = {AVAILABLE: "Available", LOW: "Low stock", OUT: "Out of stock"}

C_NAME, C_CODE, C_CAT, C_QTY, C_UNIT, C_MIN, C_STATUS, C_EXP, C_HIST = (
    "ชื่อสินค้า", "รหัส", "หมวด", "จำนวนในสต็อก", "หน่วย", "ขั้นต่ำ", "สถานะ", "วันหมดอายุ", "ประวัติ")


def _flash(kind: str, text: str) -> None:
    """Shown by app.show_flash() after st.rerun()."""
    st.session_state.setdefault("_flash", []).append((kind, text))


def _bump() -> None:
    st.session_state["stock_ver"] = st.session_state.get("stock_ver", 0) + 1


def save(change, message: str) -> bool:
    """
    Apply `change(data) -> error text` to the latest stock.json and save it (local file + Supabase /
    GitHub data branch when configured, see storage.py). The latest version is downloaded first, so
    two people saving one after the other do not overwrite each other's entries.
    """
    cfg = storage.config(st.secrets)
    data = None
    if cfg:
        try:
            latest = storage.fetch(cfg, stock.STOCK_FILE)
            data = stock.parse(latest) if latest else None
        except Exception as e:
            st.error(f"ดึงข้อมูลสต็อกล่าสุดไม่ได้ — ยังไม่ได้บันทึก ลองใหม่อีกครั้ง ({str(e)[:150]})")
            return False
    if data is None:
        data = stock.load()
    err = change(data)
    if err:
        st.error(err)
        return False
    stock.save_local(data)
    if cfg:
        try:
            warning = storage.push(cfg, stock.STOCK_FILE, stock.to_json(data), message)
        except Exception as e:
            _flash("error", f"บันทึกไว้ชั่วคราวแล้ว แต่ส่งขึ้น {storage.where(cfg)} ไม่สำเร็จ — บันทึกอีกครั้ง "
                            f"ไม่อย่างนั้นข้อมูลจะหายเมื่อแอปรีสตาร์ท: {str(e)[:200]}")
            _bump()
            return True
        if warning:
            _flash("warning", warning)
    _flash("success", message.replace("Stock: ", "บันทึกแล้ว — ", 1))
    _bump()
    return True


MAX_LOT_LINES = 3   # the table shows at most this many lots per row (every row gets the same height)
NBSP = "\u00a0"


def lots_text(item: dict, sep: str = " · ", limit: int | None = None) -> str:
    """Expiry dates of an item, earliest first: '26/03/2027 (20) · 01/05/2027 (12)'.
    With `limit`, the last part says how many more lots there are."""
    today = stock.today().isoformat()
    parts = []
    for lot in item["lots"]:
        mark = "⚠️ " if lot["exp"] < today else ""
        parts.append(f"{mark}{fmt_date(lot['exp'])} ({fmt_num(lot['qty'])})")
    if limit and len(parts) > limit:
        parts = parts[:limit - 1] + [f"…อีก {len(parts) - limit + 1} ล็อต"]
    return sep.join(parts)


def lots_cell(item: dict) -> str:
    """Expiry column of the main table: one lot per line. The table cell ignores line breaks, so each lot is
    glued with non-breaking spaces and the column is too narrow for two lots — the wrap puts each on its own line."""
    return " ".join(p.replace(" ", NBSP) for p in lots_text(item, " · ", MAX_LOT_LINES).split(" · "))


def item_label(item: dict) -> str:
    return f"{item['code']} · {item['name']}"


# --------------------------------------------------------------------------- #
# Dialogs
# --------------------------------------------------------------------------- #
@st.dialog("➕ สินค้าเข้า", width="large")
def movement_dialog(data: dict) -> None:
    items = {it["code"]: it for it in data["items"]}
    code = st.selectbox(
        "ค้นหาสินค้า", list(items), index=None, format_func=lambda c: item_label(items[c]),
        placeholder="พิมพ์ชื่อสินค้าหรือรหัส เช่น น้ำปลา หรือ I-001", key="mv_search",
    )
    if code is None:
        st.caption("พิมพ์ชื่อหรือรหัสในช่องด้านบน แล้วเลือกสินค้าที่ต้องการ")
        return
    item = items[code]
    s = stock.status(item)
    color = COLORS[s][0]
    unit = item["unit"]
    st.markdown(
        f"**{item['name']}** ({item['code']} · {item['category']}) — ตอนนี้มี "
        f"<span style='color:{color};font-weight:600'>{fmt_num(item['qty'])} {unit}</span> "
        f"· ขั้นต่ำ {fmt_num(item['min'])} · <span style='color:{color}'>{STATUS_TH[s]}</span>",
        unsafe_allow_html=True,
    )
    if item["lots"]:
        st.caption("วันหมดอายุ: " + lots_text(item))

    c_in, c_count, c_exp, c_date, c_after = st.columns([1, 1.1, 1.2, 1.2, 1.1])
    count = c_count.number_input("จำนวนในสต็อก", min_value=0.0, value=float(item["qty"]), step=1.0,
                                 format="%g", key=f"mv_count_{code}",
                                 help="ยอดตามสต็อกตอนนี้ — ถ้านับจริงได้ไม่ตรง แก้ตัวเลขนี้เป็นยอดที่นับได้")
    qty_in = c_in.number_input("เข้า (+)", min_value=0.0, value=0.0, step=1.0, format="%g", key=f"mv_in_{code}")
    exp = c_exp.date_input("วันหมดอายุ (ไม่บังคับ)", value=None, format="DD/MM/YYYY", key=f"mv_exp_{code}",
                           help="วันหมดอายุของล็อตที่รับเข้า — ไม่รู้ก็เว้นว่างได้")
    today = stock.today()
    on = c_date.date_input("วันที่บันทึก", value=today, format="DD/MM/YYYY", key=f"mv_date_{code}",
                           min_value=today - timedelta(days=stock.HISTORY_DAYS), max_value=today,
                           help=f"เก็บประวัติย้อนหลัง {stock.HISTORY_DAYS} วัน")
    after = stock.num(count + qty_in)
    after_color = COLORS[stock.status({**item, "qty": after})][0]
    c_after.markdown(
        f"<div style='font-size:0.875rem'>ยอดหลังบันทึก</div>"
        f"<div style='font-size:1.6rem;font-weight:700;color:{after_color};line-height:2.2rem'>"
        f"{fmt_num(after)} <span style='font-size:0.9rem'>{unit}</span></div>",
        unsafe_allow_html=True,
    )
    st.caption("เอาสินค้าออก: ใช้ปุ่ม 📤 เบิกสินค้าออก")

    recount = stock.num(count) != item["qty"]
    if c_after.button("💾 บันทึก", type="primary", key=f"mv_save_{code}", width="stretch"):
        def change(d: dict) -> str:
            return stock.record(d, code, qty_in, 0, count if recount else None, exp, on)
        parts = [f"+{fmt_num(qty_in)}" if qty_in else "", f"นับใหม่ {fmt_num(count)}" if recount else ""]
        msg = f"Stock: {item['name']} ({code}) {' '.join(p for p in parts if p)} → {fmt_num(after)} {unit}"
        if save(change, msg.replace("  ", " ")):
            st.session_state.pop("mv_search", None)
            st.rerun()


WEEKDAYS = ("จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์")
W_CODE, W_NAME, W_LOT, W_QTY, W_UNIT = "รหัส", "สินค้า", "ล็อต (วันหมดอายุ)", "จำนวน", "หน่วย"


def day_label(iso: str) -> str:
    """'2026-10-14' -> '14/10/2026 (พุธ)'."""
    d = date.fromisoformat(iso)
    return f"{d:%d/%m/%Y} ({WEEKDAYS[d.weekday()]})"


def lot_label(exp: str) -> str:
    return fmt_date(exp) if exp else "ไม่รู้วันหมดอายุ"


def _close_withdraw() -> None:
    for k in ("wd_keep", "wd_open", "wd_draft", "wd_msg"):
        st.session_state.pop(k, None)


def _new_list_version() -> None:
    st.session_state["wd_ver"] = st.session_state.get("wd_ver", 0) + 1


@st.dialog("📤 เบิกสินค้าออก", width="large", on_dismiss=_close_withdraw)
def withdraw_dialog() -> None:
    data = stock.load()          # read every time: the dialog stays open after 💾 บันทึกใบเบิก
    opened = st.session_state.get("wd_open")       # (delivery date, place) or None
    if opened is None:
        sheet_list(data)
    else:
        sheet_editor(data, *opened)


def _open_sheet(on: str | None, place: str = "") -> None:
    st.session_state.update(wd_open=(on, place) if on else None, wd_draft=None)
    _new_list_version()
    st.rerun(scope="fragment")


def sheet_list(data: dict) -> None:
    """Step 1: pick the delivery date + place — a new ใบเบิก, or open one made before (like a folder)."""
    st.caption("1 วันที่ส่งของ + 1 สถานที่ = 1 ใบเบิก • เลือกวันที่และสถานที่ส่ง แล้วเพิ่มสินค้าที่จะเบิกลงในรายการ • "
               "ใบเบิกที่สร้างไว้แล้ว กดเปิดเพื่อแก้ไขได้")
    today = stock.today()
    c_day, c_place, c_go = st.columns([1.1, 1.4, 1], vertical_alignment="bottom")
    day = c_day.date_input("วันที่ส่งของ (ไม่ใช่วันที่บันทึก)", value=today, format="DD/MM/YYYY", key="wd_day",
                           min_value=today - timedelta(days=stock.HISTORY_DAYS))
    place = c_place.selectbox("สถานที่ส่ง", stock.PLACES, index=None, placeholder="เลือกสถานที่ส่ง",
                              key="wd_place")
    exists = place is not None and stock.sheet(data, day.isoformat(), place) is not None
    if c_go.button("📂 เปิดใบเบิกนี้" if exists else "➕ สร้างใบเบิก", type="primary", width="stretch",
                   key="wd_go", disabled=place is None, help="เลือกสถานที่ส่งก่อน" if place is None else None):
        _open_sheet(day.isoformat(), place)

    order = {p: i for i, p in enumerate(stock.PLACES)}
    sheets = sorted(data["withdrawals"], key=lambda w: (w["date"], -order.get(w["place"], -1)), reverse=True)
    st.markdown("**📁 ใบเบิกที่มีอยู่**")
    if not sheets:
        st.caption("ยังไม่มีใบเบิก")
        return
    items = {it["code"]: it for it in data["items"]}
    box = st.container(height=min(20 + 66 * len(sheets), 400), border=True)
    for w in sheets:
        names = list(dict.fromkeys(items[l["code"]]["name"] if l["code"] in items else l["name"] for l in w["lines"]))
        more = f" …อีก {len(names) - 3}" if len(names) > 3 else ""
        c_txt, c_btn = box.columns([5, 1], vertical_alignment="center")
        c_txt.markdown(f"📁 **{day_label(w['date'])} · {stock.place_label(w['place'])}** — "
                       f"{len(w['lines'])} รายการ  \n"
                       f"<span style='font-size:0.8rem;color:#777'>{', '.join(names[:3])}{more} "
                       f"· บันทึกล่าสุด {w['saved_at']}</span>", unsafe_allow_html=True)
        if c_btn.button("เปิด", key=f"wd_open_{w['date']}_{w['place']}", width="stretch"):
            _open_sheet(w["date"], w["place"])


def _current_lines(draft: list[dict], key: str) -> list[dict]:
    """The draft with what was changed in the list table (deleted rows / new จำนวน)."""
    state = st.session_state.get(key) or {}
    deleted = {int(i) for i in state.get("deleted_rows", [])}
    edits = {int(i): e for i, e in state.get("edited_rows", {}).items()}
    out = []
    for i, line in enumerate(draft):
        if i in deleted:
            continue
        qty = edits.get(i, {}).get(W_QTY, line["qty"])
        out.append({**line, "qty": stock.num(qty) if qty is not None else 0})
    return out


def _same(a: list[dict], b: list[dict]) -> bool:
    return [(l["code"], l["exp"], stock.num(l["qty"])) for l in a] == \
           [(l["code"], l["exp"], stock.num(l["qty"])) for l in b]


def sheet_excel(data: dict, on: str, place: str) -> bytes:
    """The saved ใบเบิก of one delivery date + place as an .xlsx file (A4, fits the page width when printed)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    w = stock.sheet(data, on, place)
    items = {it["code"]: it for it in data["items"]}
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="2F5597")
    stripe = PatternFill("solid", fgColor="DDEBF7")
    center = Alignment(horizontal="center", vertical="center")
    cols = [("ลำดับ", 7), ("รหัส", 9), ("สินค้า", 34), ("หมวด", 13), ("วันหมดอายุ (ล็อต)", 18),
            ("จำนวนเบิก", 12), ("หน่วย", 9), ("หมายเหตุ", 22)]

    wb = Workbook()
    ws = wb.active
    ws.title = f"ใบเบิก {date.fromisoformat(on):%d-%m-%Y}"
    ws["A1"] = "ใบเบิกสินค้า"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = f"วันที่ส่งของ: {day_label(on)}    สถานที่ส่ง: {stock.place_label(place)}"
    ws["A2"].font = Font(size=12, bold=True)
    ws["A3"] = f"บันทึกล่าสุด: {w['saved_at']}  •  {len(w['lines'])} รายการ"
    ws["A3"].font = Font(size=10, color="666666")
    top = 5
    for c, (title, width) in enumerate(cols, 1):
        cell = ws.cell(top, c, title)
        cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, color="FFFFFF"), head_fill, center, box
        ws.column_dimensions[cell.column_letter].width = width
    for i, line in enumerate(w["lines"], 1):
        it = items.get(line["code"])
        row = [i, line["code"], it["name"] if it else f"{line['name']} (ลบแล้ว)",
               it["category"] if it else "", lot_label(line["exp"]), line["qty"],
               it["unit"] if it else line.get("unit", ""), ""]
        r = top + i
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.border = box
            if i % 2 == 0:
                cell.fill = stripe
            if c in (1, 2, 5, 7):
                cell.alignment = center
    last = top + len(w["lines"])
    ws.cell(last + 1, 5, "รวม").font = Font(bold=True)
    ws.cell(last + 1, 5).alignment = Alignment(horizontal="right")
    ws.cell(last + 1, 6, f"{len(w['lines'])} รายการ").font = Font(bold=True)
    ws.cell(last + 3, 2, "ผู้เบิก ..................................")
    ws.cell(last + 3, 5, "ผู้จ่าย ..................................")
    ws.freeze_panes = ws.cell(top + 1, 1)
    ws.print_title_rows = f"{top}:{top}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheet_editor(data: dict, on: str, place: str) -> None:
    """Step 2: the list of one ใบเบิก — add item + lot + จำนวน, edit / delete lines, save (cuts the stock)."""
    saved = stock.sheet(data, on, place)
    saved_lines = [{"code": l["code"], "exp": l["exp"], "qty": l["qty"]} for l in (saved or {"lines": []})["lines"]]
    old_names = {l["code"]: l for l in (saved or {"lines": []})["lines"]}
    if st.session_state.get("wd_draft") is None:
        st.session_state["wd_draft"] = [dict(l) for l in saved_lines]
    draft = st.session_state["wd_draft"]
    ver = st.session_state.get("wd_ver", 0)
    list_key = f"wd_list_{ver}"
    lines = _current_lines(draft, list_key)
    items = {it["code"]: it for it in data["items"]}

    c_back, c_title = st.columns([1, 3.2], vertical_alignment="center")
    if c_back.button("⬅️ ใบเบิกทั้งหมด", key="wd_back", width="stretch"):
        _open_sheet(None)
    c_title.markdown(f"#### 📁 ใบเบิก ส่งวันที่ {day_label(on)}  \n📍 {stock.place_label(place)}")
    msg = st.session_state.pop("wd_msg", None)
    if msg:
        st.success(msg)
    st.caption((f"บันทึกล่าสุด {saved['saved_at']}" if saved else "ใบเบิกใหม่ — ยังไม่ได้บันทึก") +
               " • สต็อกจะถูกตัดตอนกด 💾 บันทึกใบเบิก • แก้ทีหลังได้ ระบบปรับสต็อกตามส่วนต่าง")

    # ---- add a line: item -> lot (earliest expiry first) -> จำนวน
    st.markdown("**➕ เพิ่มสินค้าที่จะเบิก**")
    c_item, c_lot, c_qty, c_add = st.columns([2.2, 2.3, 1, 0.8], vertical_alignment="bottom")
    code = c_item.selectbox("สินค้า", list(items), index=None, format_func=lambda c: item_label(items[c]),
                            placeholder="พิมพ์ชื่อสินค้าหรือรหัส", key=f"wd_item_{ver}")
    lot, qty, left = None, 0.0, {}
    if code:
        item = items[code]
        unit = item["unit"]
        used: dict[str, float] = {}
        for l in lines:
            if l["code"] == code:
                used[l["exp"]] = stock.num(used.get(l["exp"], 0) + l["qty"])
        have = stock.lots_for_sheet(data, on, code, place)
        left = {e: stock.num(q - used.get(e, 0)) for e, q in have.items()}
        options = [e for e in sorted(left) if e and left[e] > 0] + ([""] if left.get("", 0) > 0 else [])
        today = stock.today().isoformat()

        def label(e: str) -> str:
            when = ("⚠️ หมดอายุแล้ว " if e and e < today else "หมดอายุ ") + fmt_date(e) if e else "ไม่รู้วันหมดอายุ"
            first = " (หมดก่อน)" if e and e == options[0] else ""
            return f"{when} — เหลือ {fmt_num(left[e])} {unit}".rstrip() + first

        if options:
            lot = c_lot.selectbox("ล็อต (วันหมดอายุ)", options, format_func=label, key=f"wd_lot_{ver}_{code}",
                                  help="เลือกล็อตที่จะเบิก — ตั้งไว้ที่ล็อตที่หมดอายุก่อน")
            qty = c_qty.number_input(f"จำนวน {unit}".strip(), min_value=0.0, max_value=float(left[lot]),
                                     value=float(min(1, left[lot])), step=1.0, format="%g",
                                     key=f"wd_qty_{ver}_{code}_{lot}")
        else:
            c_lot.selectbox("ล็อต (วันหมดอายุ)", [], placeholder="หมดสต็อก — ไม่มีของให้เบิก", disabled=True,
                            key=f"wd_lot_none_{ver}")
            c_qty.number_input("จำนวน", value=0.0, disabled=True, key=f"wd_qty_none_{ver}")
    else:
        c_lot.selectbox("ล็อต (วันหมดอายุ)", [], placeholder="เลือกสินค้าก่อน", disabled=True,
                        key=f"wd_lot_none_{ver}")
        c_qty.number_input("จำนวน", value=0.0, disabled=True, key=f"wd_qty_none_{ver}")
    if c_add.button("➕ เพิ่ม", key=f"wd_add_{ver}", disabled=lot is None or qty <= 0, width="stretch"):
        new = [dict(l) for l in lines]
        same = next((l for l in new if l["code"] == code and l["exp"] == lot), None)
        if same:
            same["qty"] = stock.num(same["qty"] + qty)
        else:
            new.append({"code": code, "exp": lot, "qty": stock.num(qty)})
        st.session_state["wd_draft"] = new
        _new_list_version()
        st.rerun(scope="fragment")

    # ---- the list
    st.markdown(f"**📋 รายการเบิก** ({len(lines)} รายการ)")
    if not draft:
        st.info("ยังไม่มีสินค้าในใบเบิกนี้ — เลือกสินค้าด้านบนแล้วกด ➕ เพิ่ม")
    else:
        def name_of(c: str) -> str:
            return items[c]["name"] if c in items else f"{old_names.get(c, {}).get('name', '')} (ลบแล้ว)".lstrip()
        df = pd.DataFrame([{
            W_CODE: l["code"], W_NAME: name_of(l["code"]), W_LOT: lot_label(l["exp"]), W_QTY: float(l["qty"]),
            W_UNIT: items[l["code"]]["unit"] if l["code"] in items else old_names.get(l["code"], {}).get("unit", ""),
        } for l in draft])
        st.data_editor(
            df, num_rows="delete", hide_index=True, width="stretch", key=list_key,
            disabled=[W_CODE, W_NAME, W_LOT, W_UNIT],
            column_config={
                W_CODE: st.column_config.TextColumn(width=70),
                W_NAME: st.column_config.TextColumn(width=230),
                W_LOT: st.column_config.TextColumn(width=140),
                W_QTY: st.column_config.NumberColumn(min_value=0, required=True, format="localized", width=90,
                                                     help="แก้จำนวนได้ในตาราง"),
                W_UNIT: st.column_config.TextColumn(width=70),
            },
        )
        st.caption("แก้จำนวน: กดช่องจำนวน • ลบรายการ: ติ๊กหน้าแถวแล้วกดถังขยะมุมขวาบนของตาราง")

    # ---- check every item + lot against what is in stock
    errors = []
    totals: dict[tuple[str, str], float] = {}
    for l in lines:
        totals[(l["code"], l["exp"])] = stock.num(totals.get((l["code"], l["exp"]), 0) + l["qty"])
    for (c, e), q in totals.items():
        if c not in items:
            errors.append(f"{old_names.get(c, {}).get('name', c)} ถูกลบออกจากสต็อกแล้ว — ลบแถวนี้ออก")
            continue
        have = stock.lots_for_sheet(data, on, c, place).get(e, 0)
        if q > have:
            errors.append(f"{items[c]['name']} ล็อต {lot_label(e)} มีแค่ {fmt_num(have)} {items[c]['unit']} "
                          f"(ในรายการ {fmt_num(q)})")
    for err in errors:
        st.error(err)

    changed = not _same(lines, saved_lines)
    s1, s_xl, s2, s3 = st.columns([1.3, 1.4, 1.3, 1.6], vertical_alignment="center")
    if s1.button("💾 บันทึกใบเบิก", type="primary", key=f"wd_save_{ver}", width="stretch",
                 disabled=not changed or bool(errors) or (not lines and not saved)):
        total = sum(1 for l in lines if l["qty"] > 0)
        where = f"ส่ง {fmt_date(on)} {stock.place_label(place)}"
        what = f"ใบเบิก {where} — {total} รายการ" if total else f"ลบใบเบิก {where}"
        if save(lambda d: stock.save_withdrawal(d, date.fromisoformat(on), lines, place), f"Stock: {what}"):
            st.session_state.update(wd_draft=None, wd_msg="✅ บันทึกใบเบิกแล้ว — ตัดสต็อกเรียบร้อย",
                                    wd_open=(on, place) if total else None)
            _new_list_version()
            st.rerun()                   # whole page: the stock table behind shows the new numbers
    s_xl.download_button(
        "📥 Download Excel", sheet_excel(data, on, place) if saved else b"", key=f"wd_xlsx_{ver}",
        width="stretch", file_name=f"ใบเบิก_{on}_{stock.place_label(place).replace(' ', '_')}.xlsx", on_click="ignore", disabled=not saved,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help=("ไฟล์นี้เป็นใบเบิกที่บันทึกล่าสุด — กด 💾 บันทึกใบเบิก ก่อน ถ้าอยากได้รายการที่เพิ่งแก้"
              if changed else "รายการที่เบิกในวันนี้เป็นไฟล์ Excel") if saved else "บันทึกใบเบิกก่อนจึงดาวน์โหลดได้")
    if changed:
        if s2.button("↩️ ยกเลิกแก้ไข", key=f"wd_reset_{ver}", width="stretch"):
            st.session_state["wd_draft"] = None
            _new_list_version()
            st.rerun(scope="fragment")
        st.caption("⚠️ มีการแก้ไขที่ยังไม่ได้บันทึก — สต็อกยังไม่ถูกตัด (Excel ยังเป็นใบเบิกที่บันทึกล่าสุด)")
    elif saved:
        with s2.popover("🗑️ ลบใบเบิกนี้", width="stretch"):
            st.caption("ลบแล้วสินค้าในใบเบิกนี้จะกลับเข้าสต็อก (เข้าล็อตเดิม)")
            if st.button("ยืนยันลบใบเบิก", type="primary", key=f"wd_del_{ver}"):
                if save(lambda d: stock.save_withdrawal(d, date.fromisoformat(on), [], place),
                        f"Stock: ลบใบเบิก ส่ง {fmt_date(on)} {stock.place_label(place)}"):
                    st.session_state.update(wd_draft=None, wd_open=None,
                                            wd_msg=None)
                    _new_list_version()
                    st.rerun()


@st.dialog("🗑️ ลบสินค้า")
def delete_dialog(data: dict) -> None:
    items = {it["code"]: it for it in data["items"]}
    code = st.selectbox(
        "ค้นหาสินค้าที่จะลบ", list(items), index=None, format_func=lambda c: item_label(items[c]),
        placeholder="พิมพ์ชื่อสินค้าหรือรหัส", key="del_search",
    )
    if code is None:
        st.caption("พิมพ์ชื่อหรือรหัสในช่องด้านบน แล้วเลือกสินค้าที่ต้องการลบ")
        return
    item = items[code]
    unit = item["unit"]
    st.markdown(f"**{item['name']}** ({item['code']} · {item['category']}) — ตอนนี้มี "
                f"**{fmt_num(item['qty'])} {unit}**".rstrip())
    if item["lots"]:
        st.caption("วันหมดอายุ: " + lots_text(item))
    if item["qty"] > 0:
        st.warning(f"สินค้านี้ยังมีในสต็อก {fmt_num(item['qty'])} {unit} — ลบแล้วยอดนี้จะหายไปด้วย".replace("  ", " "))
    st.caption("ประวัติการเข้า-ออกเดิมยังดูได้ใน 📜 ประวัติทั้งหมด (ขึ้นว่า ลบแล้ว) • ลบแล้วกู้คืนไม่ได้ "
               "ต้องเพิ่มสินค้าใหม่")
    sure = st.checkbox(f"ยืนยันว่าจะลบ {item['name']}", key=f"del_sure_{code}")
    if st.button("🗑️ ลบสินค้า", type="primary", disabled=not sure, key=f"del_go_{code}"):
        if save(lambda d: stock.delete_item(d, code), f"Stock: ลบสินค้า {item['name']} ({code})"):
            st.session_state.pop("del_search", None)
            st.rerun()


@st.dialog("🆕 เพิ่มสินค้าใหม่")
def add_dialog(data: dict) -> None:
    st.caption("ใส่แค่ชื่อ รหัส และหมวด — จำนวน ขั้นต่ำ วันหมดอายุ ไปกรอกตอนสินค้าเข้า / ในตารางหลัก")
    name = st.text_input("ชื่อสินค้า", key="add_name")
    cat = st.selectbox("หมวด", list(CATEGORIES), index=None, placeholder="เลือกหมวด",
                       format_func=lambda c: f"{c} ({CATEGORIES[c]})", key="add_cat")
    code = st.text_input("รหัส", value=stock.next_code(data, cat) if cat else "", key=f"add_code_{cat}",
                         disabled=cat is None, help="ตั้งให้อัตโนมัติ = ตัวอักษรหมวด + เลขถัดไป (แก้ได้)")
    unit = st.text_input("หน่วย (ไม่บังคับ)", key="add_unit", placeholder="เช่น กก., ขวด, แพ็ค")
    if st.button("💾 เพิ่มสินค้า", type="primary", disabled=not (name.strip() and cat and code.strip())):
        if save(lambda d: stock.add_item(d, name, code, cat, unit),
                f"Stock: เพิ่มสินค้าใหม่ {name.strip()} ({code.strip().upper()})"):
            for k in ("add_name", "add_cat", "add_unit"):
                st.session_state.pop(k, None)
            st.rerun()


# --------------------------------------------------------------------------- #
# Page
# --------------------------------------------------------------------------- #
def summary_cards(items: list[dict]) -> None:
    counts = {s: sum(stock.status(it) == s for it in items) for s in COLORS}
    cards = [("ทั้งหมด", len(items), "#333", "#999")] + [
        (STATUS_TH[s], counts[s], COLORS[s][0], COLORS[s][0]) for s in COLORS]
    html = "".join(
        f"<div style='flex:1;min-width:130px;border:1px solid #ddd;border-left:5px solid {edge};"
        f"border-radius:10px;padding:10px 16px'>"
        f"<span style='font-size:1.9rem;font-weight:700;color:{color}'>{n}</span> "
        f"<span style='font-size:0.85rem;letter-spacing:0.05em;color:{color}'>{label.upper()}</span></div>"
        for label, n, color, edge in cards)
    st.markdown(f"<div style='display:flex;gap:12px;flex-wrap:wrap;margin:4px 0 12px'>{html}</div>",
                unsafe_allow_html=True)


def render(show_flash=None) -> None:
    st.title("📦 สต็อกสินค้า")
    if show_flash:
        show_flash()
    st.caption("ดูยอดคงเหลือของสินค้าทั้งหมด • กด **📜 ดูประวัติ** ท้ายแถวเพื่อดูการเข้า-ออกของสินค้านั้น • "
               "กดช่อง **วันหมดอายุ** เพื่อแก้วันหมดอายุ • "
               "เปิด **✏️ แก้ไขตาราง** เพื่อแก้ ชื่อสินค้า / หมวด / หน่วย / ขั้นต่ำ "
               "(ย้ายหมวด = ได้รหัสใหม่เป็นเลขถัดไปของหมวดนั้น)")
    if not storage.config(st.secrets):
        st.info("💡 ตอนนี้บันทึกลงไฟล์ในเครื่องที่รันแอปเท่านั้น — ถ้าแอปอยู่บน Streamlit Cloud "
                "ข้อมูลจะหายเมื่อแอปรีสตาร์ท จนกว่าจะตั้งค่า Supabase หรือ GitHub ใน Secrets (ดู README)")
    data = stock.load()
    items = data["items"]
    ver = st.session_state.get("stock_ver", 0)
    dirty = st.session_state.get("stock_dirty", False)

    b1, b0, b2, b3, b4, _ = st.columns([1.1, 1.2, 1.2, 1.2, 1.1, 0.6])
    if b1.button("➕ สินค้าเข้า", type="primary", width="stretch", disabled=dirty):
        st.session_state["mv_search"] = None
        movement_dialog(data)
    if b0.button("📤 เบิกสินค้าออก", type="primary", width="stretch", disabled=dirty):
        st.session_state.update(wd_keep=True, wd_open=None, wd_draft=None)
    if st.session_state.get("wd_keep"):          # stays open after 💾 บันทึกใบเบิก until closed with ✕
        withdraw_dialog()
    if b2.button("🆕 เพิ่มสินค้าใหม่", width="stretch", disabled=dirty):
        add_dialog(data)
    if b3.button("📜 ประวัติทั้งหมด", width="stretch", disabled=dirty):
        all_history_dialog(data)
    if b4.button("🗑️ ลบสินค้า", width="stretch", disabled=dirty or not items):
        st.session_state["del_search"] = None
        delete_dialog(data)
    open_code = st.session_state.pop("hist_open", None)   # set by a click in the table, see below
    if open_code and stock.find(data, open_code):
        history_dialog(data, open_code)
    exp_code = st.session_state.pop("exp_open", None)
    if exp_code and stock.find(data, exp_code):
        expiry_dialog(data, exp_code)

    f1, f2 = st.columns([3, 1])
    query = f1.text_input("ค้นหา", placeholder="ค้นหาชื่อสินค้าหรือรหัส...", label_visibility="collapsed",
                          key="stock_query", disabled=dirty)
    cat_filter = f2.selectbox("หมวด", ["ทุกหมวด"] + list(CATEGORIES), label_visibility="collapsed",
                              key="stock_cat", disabled=dirty)
    if dirty:
        st.caption("⚠️ บันทึกหรือยกเลิกการแก้ไขในตารางก่อน จึงจะค้นหา / รับเข้า / เบิกสินค้าได้")
    summary_cards(items)

    q = query.strip().lower()
    view = [it for it in items
            if (cat_filter == "ทุกหมวด" or it["category"] == cat_filter)
            and (not q or q in it["name"].lower() or q in it["code"].lower())]
    if not view:
        st.info("ไม่พบสินค้า" if items else "ยังไม่มีสินค้า — กด 🆕 เพิ่มสินค้าใหม่")
        return

    status_by_row = [stock.status(it) for it in view]

    def row_colors(row: pd.Series) -> list[str]:
        fg, bg = COLORS[status_by_row[row.name]]
        return [f"color:{fg};font-weight:600" if c == C_QTY else
                f"color:{fg};background-color:{bg};font-weight:600" if c == C_STATUS else "" for c in row.index]

    base_config = {
        C_NAME: st.column_config.TextColumn(width=230),
        C_CODE: st.column_config.TextColumn(width=70),
        C_CAT: st.column_config.TextColumn(width=95),
        C_QTY: st.column_config.NumberColumn(width=115, help="เปลี่ยนยอดที่ปุ่ม ➕ สินค้าเข้า / 📤 เบิกสินค้าออก"),
        C_UNIT: st.column_config.TextColumn(width=70),
        C_MIN: st.column_config.NumberColumn(width=70, help="ต่ำกว่าขั้นต่ำ = Low stock • 0 = Out of stock"),
        C_STATUS: st.column_config.TextColumn(width=110),
        C_EXP: st.column_config.TextColumn(width=175, help="ทุกล็อตที่ยังเหลือ (จำนวน) • ⚠️ = หมดอายุแล้ว "
                                                               "• กดเพื่อแก้วันหมดอายุ"),
    }
    lines = min(max([len(it["lots"]) for it in view] + [1]), MAX_LOT_LINES)   # one lot per line
    row_height = 35 if lines == 1 else 12 + 21 * lines
    height = min(38 + row_height * len(view), 700)
    if st.session_state.pop("stock_edit_off", False):
        st.session_state["stock_edit"] = False
    editing = st.toggle("✏️ แก้ไขตาราง (ชื่อสินค้า / หมวด / หน่วย / ขั้นต่ำ)", key="stock_edit", disabled=dirty,
                        help="บันทึกหรือยกเลิกการแก้ไขก่อนปิด" if dirty else None)

    if not editing:
        # ---- view: click 📜 ดูประวัติ in a row -> that item's history pop-up
        df = pd.DataFrame([{
            C_NAME: it["name"], C_CODE: it["code"], C_CAT: it["category"], C_QTY: it["qty"], C_UNIT: it["unit"],
            C_MIN: it["min"], C_STATUS: STATUS_TH[stock.status(it)], C_EXP: lots_cell(it), C_HIST: "📜 ดูประวัติ",
        } for it in view])
        event = st.dataframe(
            df.style.apply(row_colors, axis=1).format({C_QTY: fmt_num, C_MIN: fmt_num}),
            hide_index=True, width="content", height=height, row_height=row_height,
            column_config={**base_config,
                           C_HIST: st.column_config.TextColumn(width=110, help="กดเพื่อดูประวัติการเข้า-ออก")},
            on_select="rerun", selection_mode="single-cell",
            key=f"stock_table_{st.session_state.get('hist_ver', 0)}_{cat_filter}_{q}",
        )
        cells = event.selection.cells if event else []
        if cells and cells[0][1] in (C_HIST, C_EXP) and 0 <= cells[0][0] < len(view):
            target = "hist_open" if cells[0][1] == C_HIST else "exp_open"
            st.session_state[target] = view[cells[0][0]]["code"]
            st.session_state["hist_ver"] = st.session_state.get("hist_ver", 0) + 1   # clear the selection
            st.rerun()
        return

    # ---- edit: name / หมวด / หน่วย / ขั้นต่ำ
    df = pd.DataFrame([{
        C_NAME: it["name"], C_CODE: it["code"], C_CAT: it["category"], C_QTY: it["qty"], C_UNIT: it["unit"],
        C_MIN: it["min"], C_STATUS: STATUS_TH[stock.status(it)], C_EXP: lots_cell(it),
    } for it in view])
    edited = st.data_editor(
        df.style.apply(row_colors, axis=1).format({C_QTY: fmt_num, C_MIN: fmt_num}),
        hide_index=True,
        width="content",
        height=height,
        row_height=row_height,
        disabled=[C_CODE, C_QTY, C_STATUS, C_EXP],
        column_config={
            **base_config,
            C_NAME: st.column_config.TextColumn(required=True, width=230),
            C_CAT: st.column_config.SelectboxColumn(options=list(CATEGORIES), required=True, width=110,
                                                    help="ย้ายหมวด = ได้รหัสใหม่ (ตัวอักษรหมวด + เลขถัดไป) ตอนบันทึก"),
            C_MIN: st.column_config.NumberColumn(min_value=0, width=70, format="localized",
                                                 help="ต่ำกว่าขั้นต่ำ = Low stock • 0 = Out of stock"),
        },
        key=f"stock_editor_{ver}_{cat_filter}_{q}",
    )
    changes = []
    for it, (_, row) in zip(view, edited.iterrows()):
        new = (str(row[C_NAME] or "").strip(), row[C_CAT], str(row[C_UNIT] or "").strip(),
               stock.num(0 if pd.isna(row[C_MIN]) else row[C_MIN]))
        if new != (it["name"], it["category"], it["unit"], it["min"]):
            changes.append((it["code"], *new))
    if bool(changes) != dirty:
        st.session_state["stock_dirty"] = bool(changes)
        st.rerun()
    if changes:
        s1, s2, s3 = st.columns([1, 1, 3])
        if s1.button("💾 บันทึก", type="primary", key="stock_save"):
            moved = []

            def change(d: dict) -> str:
                for code, name, cat, unit, min_level in changes:
                    err, new_code = stock.update_item(d, code, name, cat, unit, min_level)
                    if err:
                        return f"{code}: {err}"
                    if new_code != code:
                        moved.append(f"{name}: {code} → {new_code}")
                return ""
            if save(change, f"Stock: แก้ไขข้อมูลสินค้า {len(changes)} รายการ"):
                for m in moved:
                    _flash("info", f"ย้ายหมวดแล้ว รหัสใหม่ — {m}")
                st.session_state["stock_dirty"] = False
                st.session_state["stock_edit_off"] = True
                st.rerun()
        if s2.button("↩️ ยกเลิกการแก้ไข", key="stock_reset"):
            st.session_state["stock_dirty"] = False
            st.session_state["stock_edit_off"] = True
            _bump()
            st.rerun()
        s3.caption(f"⚠️ แก้ไข {len(changes)} รายการ ยังไม่ได้บันทึก")


def history_rows(rows: list[dict], items: dict[str, dict] | None = None, unit: str = "") -> pd.DataFrame:
    """Movements as a table; with `items` (code -> item) the product is shown too (all-items history)."""
    out = []
    for m in rows:
        it = (items or {}).get(m["code"])
        u = it["unit"] if it else m.get("unit", unit)
        row = {"วันที่": fmt_date(m["date"])}
        if items is not None:
            row.update({"รหัส": m["code"], "สินค้า": it["name"] if it else f"{m.get('name', '')} (ลบแล้ว)".lstrip()})
        row.update({
            "เข้า (+)": fmt_num(m["in"]) if m["in"] else "",
            "ออก (−)": fmt_num(m["out"]) if m["out"] else "",
            "ยอดก่อน": fmt_num(m["before"]),
            "นับใหม่": fmt_num(m["count"]) if m["count"] != m["before"] else "",
            "ยอดหลัง": f"{fmt_num(m['after'])} {u}".strip(),
            "วันหมดอายุ": " · ".join(
                ([f"เข้า {fmt_date(m['exp'])}"] if m["exp"] and m["in"] else [fmt_date(m["exp"])] if m["exp"] else [])
                + [f"ออก {fmt_date(t['exp']) if t['exp'] else 'ไม่รู้วันหมดอายุ'} ({fmt_num(t['qty'])})"
                   for t in m.get("out_lots", [])]
                + [f"คืน {fmt_date(t['exp']) if t['exp'] else 'ไม่รู้วันหมดอายุ'} ({fmt_num(t['qty'])})"
                   for t in m.get("back_lots", [])]),
            "หมายเหตุ": m.get("note", ""),
            "บันทึกเมื่อ": m.get("saved_at", ""),
        })
        out.append(row)
    return pd.DataFrame(out)


@st.dialog("📜 ประวัติทั้งหมด", width="large")
def all_history_dialog(data: dict) -> None:
    items = {it["code"]: it for it in data["items"]}
    rows = sorted(enumerate(data["movements"]),
                  key=lambda r: (r[1]["date"], r[1].get("saved_at", ""), r[0]), reverse=True)
    rows = [m for _, m in rows]
    if not rows:
        st.caption(f"ยังไม่มีการเข้า-ออกใน {stock.HISTORY_DAYS} วันที่ผ่านมา")
        return
    f1, f2 = st.columns([3, 1])
    q = f1.text_input("ค้นหา", placeholder="ค้นหาชื่อสินค้าหรือรหัส...", label_visibility="collapsed",
                      key="all_hist_q").strip().lower()
    cat = f2.selectbox("หมวด", ["ทุกหมวด"] + list(CATEGORIES), label_visibility="collapsed", key="all_hist_cat")
    if q or cat != "ทุกหมวด":
        def keep(m: dict) -> bool:
            it = items.get(m["code"], {})
            return ((cat == "ทุกหมวด" or it.get("category") == cat)
                    and (not q or q in m["code"].lower() or q in it.get("name", "").lower()))
        rows = [m for m in rows if keep(m)]
    st.caption(f"{len(rows)} รายการ ย้อนหลัง {stock.HISTORY_DAYS} วัน (ใหม่สุดอยู่บน)")
    if rows:
        st.dataframe(history_rows(rows, items), hide_index=True, width="stretch",
                     height=min(38 + 35 * len(rows), 560))


@st.dialog("📜 ประวัติการเข้า-ออก", width="large")
def history_dialog(data: dict, code: str) -> None:
    item = stock.find(data, code)
    last = stock.last_movement(data).get(code)
    unit = item["unit"]
    s = stock.status(item)
    st.markdown(
        f"**{item['name']}** ({item['code']} · {item['category']}) — ตอนนี้มี "
        f"<span style='color:{COLORS[s][0]};font-weight:600'>{fmt_num(item['qty'])} {unit}</span> "
        f"· <span style='color:{COLORS[s][0]}'>{STATUS_TH[s]}</span>",
        unsafe_allow_html=True,
    )
    rows = stock.history(data, item["code"])
    if not rows:
        st.caption(f"ยังไม่มีการเข้า-ออกใน {stock.HISTORY_DAYS} วันที่ผ่านมา")
    else:
        st.caption(f"{len(rows)} รายการ ย้อนหลัง {stock.HISTORY_DAYS} วัน • ล่าสุด {fmt_date(last)}")
        st.dataframe(history_rows(rows, unit=unit), hide_index=True, width="stretch")

    st.caption("แก้วันหมดอายุ: กดช่อง วันหมดอายุ ของสินค้านี้ในตาราง")


@st.dialog("🗓️ แก้วันหมดอายุ", width="large")
def expiry_dialog(data: dict, code: str) -> None:
    item = stock.find(data, code)
    unit = item["unit"]
    dated = sum(l["qty"] for l in item["lots"])
    st.markdown(f"**{item['name']}** ({item['code']}) — ในสต็อก **{fmt_num(item['qty'])} {unit}**".rstrip())
    st.caption(f"มีวันหมดอายุ {fmt_num(dated)} • ไม่รู้วันหมดอายุ {fmt_num(stock.num(item['qty'] - dated))} — "
               "รวมทุกล็อตต้องไม่เกินยอดในสต็อก • แก้วันที่ / จำนวนในตารางได้เลย • "
               "เพิ่มล็อต: ปุ่ม + ใต้ตาราง • ลบล็อต: ติ๊กหน้าแถวแล้วกดถังขยะ")
    lots_df = pd.DataFrame({"วันหมดอายุ": pd.to_datetime([l["exp"] for l in item["lots"]]).date,
                            "จำนวน": [float(l["qty"]) for l in item["lots"]]})
    ver = st.session_state.get("stock_ver", 0)
    lots = st.data_editor(
        lots_df, num_rows="dynamic", hide_index=True, width="stretch", key=f"lots_{code}_{ver}",
        column_config={
            "วันหมดอายุ": st.column_config.DateColumn(format="DD/MM/YYYY", required=True),
            "จำนวน": st.column_config.NumberColumn(min_value=0, required=True, format="localized"),
        },
    )
    new_lots = [(pd.Timestamp(d).date(), float(q)) for d, q in zip(lots["วันหมดอายุ"], lots["จำนวน"])
                if not pd.isna(d) and not pd.isna(q)]
    current = [(pd.Timestamp(l["exp"]).date(), float(l["qty"])) for l in item["lots"]]
    total = stock.num(sum(q for _, q in new_lots))
    if total > item["qty"]:
        st.error(f"รวมทุกล็อต {fmt_num(total)} เกินยอดในสต็อก {fmt_num(item['qty'])} {unit}")
    if st.button("💾 บันทึกวันหมดอายุ", type="primary", key=f"lots_save_{code}",
                 disabled=new_lots == current or total > item["qty"]):
        if save(lambda d: stock.set_lots(d, code, new_lots),
                f"Stock: แก้วันหมดอายุ {item['name']} ({code})"):
            st.rerun()
